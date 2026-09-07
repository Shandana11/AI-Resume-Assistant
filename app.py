import json
import re
from io import BytesIO

import streamlit as st
from google import genai
from pypdf import PdfReader
from docx import Document

st.set_page_config(
    page_title="ATS Resume Checker",
    page_icon="📄",
    layout="wide",
)

MODEL_NAME = "gemini-3.5-flash"


def get_api_key():
    """Read Gemini API key from Streamlit secrets."""
    try:
        return st.secrets["GEMINI_API_KEY"]
    except Exception:
        return None


def extract_pdf_text(file_bytes):
    reader = PdfReader(BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages).strip()


def extract_docx_text(file_bytes):
    document = Document(BytesIO(file_bytes))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    # Also capture text from tables, which are common in resumes.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                paragraphs.append(" | ".join(cells))

    return "\n".join(paragraphs).strip()


def extract_resume_text(uploaded_file):
    data = uploaded_file.getvalue()
    name = uploaded_file.name.lower()

    if name.endswith(".pdf"):
        return extract_pdf_text(data)
    if name.endswith(".docx"):
        return extract_docx_text(data)

    raise ValueError("Unsupported file type. Please upload a PDF or DOCX resume.")


def heuristic_ats_score(text):
    """A quick deterministic score used as a baseline before Gemini analysis."""
    if not text.strip():
        return 0, ["No readable text was extracted from the resume."]

    lower = text.lower()
    score = 0
    strengths = []
    issues = []

    # Contact information
    contact_checks = {
        "email": r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b",
        "phone": r"(\+?\d[\d\s().-]{7,}\d)",
        "linkedin": r"linkedin\.com",
    }

    if re.search(contact_checks["email"], text, re.I):
        score += 10
        strengths.append("Email address detected.")
    else:
        issues.append("Add a professional email address.")

    if re.search(contact_checks["phone"], text, re.I):
        score += 5
        strengths.append("Phone number detected.")
    else:
        issues.append("Add a phone number if appropriate for the target market.")

    if re.search(contact_checks["linkedin"], text, re.I):
        score += 5
        strengths.append("LinkedIn profile detected.")
    else:
        issues.append("Consider adding a LinkedIn URL.")

    # Common ATS-friendly sections
    section_groups = {
        "summary": ["summary", "professional summary", "profile", "objective"],
        "experience": ["experience", "work experience", "employment"],
        "education": ["education", "academic"],
        "skills": ["skills", "technical skills", "core skills"],
    }

    for section, keywords in section_groups.items():
        if any(k in lower for k in keywords):
            score += 12
            strengths.append(f"{section.title()} section detected.")
        else:
            issues.append(f"Consider adding a clear {section.title()} section.")

    # Quantification
    numbers = re.findall(r"\b\d+(?:\.\d+)?%?\b", text)
    if len(numbers) >= 5:
        score += 8
        strengths.append("Resume contains several measurable details.")
    else:
        issues.append("Add measurable achievements (%, $, time, volume, growth, etc.).")

    # Action verbs
    action_verbs = [
        "led", "managed", "created", "developed", "designed", "improved",
        "increased", "reduced", "built", "launched", "implemented",
        "optimized", "analyzed", "delivered", "coordinated", "achieved",
    ]
    verb_hits = sum(1 for verb in action_verbs if re.search(rf"\b{verb}\b", lower))
    if verb_hits >= 4:
        score += 10
        strengths.append("Uses several achievement/action verbs.")
    else:
        issues.append("Use stronger action verbs and achievement-focused bullets.")

    # Length/readability baseline
    word_count = len(re.findall(r"\b[\w'-]+\b", text))
    if 250 <= word_count <= 1200:
        score += 10
    elif word_count < 250:
        issues.append("The resume may be too short; add relevant evidence of experience.")
    else:
        issues.append("The resume may be too long; remove low-value or repetitive content.")

    # ATS-unfriendly characters / layout hints
    if "│" in text or "◆" in text or "▪" in text:
        issues.append("Avoid decorative symbols that may be parsed inconsistently by ATS software.")

    score = max(0, min(100, score))
    return score, issues[:8]


def analyze_with_gemini(resume_text, job_description=""):
    api_key = get_api_key()
    if not api_key:
        raise RuntimeError(
            "GEMINI_API_KEY is missing. Add it in Streamlit Secrets."
        )

    client = genai.Client(api_key=api_key)

    job_context = (
        job_description.strip()
        if job_description.strip()
        else "No job description was provided. Evaluate general ATS readiness."
    )

    prompt = f"""
You are an expert ATS resume reviewer and career coach.

Analyze the resume below. The goal is to estimate how well the resume is likely
to perform in an Applicant Tracking System and how well it communicates value
to a recruiter.

Important:
- Do not invent experience, employers, degrees, skills, dates, metrics, or achievements.
- Give practical improvements based only on the supplied resume.
- ATS scores are estimates, not guarantees. Different ATS platforms score resumes differently.
- Focus on parsing, structure, keywords, clarity, measurable achievements, and relevance.
- If a job description is provided, assess keyword alignment against it.

Return ONLY valid JSON with this exact structure:
{{
  "ats_score": 0,
  "score_explanation": "short explanation",
  "strengths": ["...", "..."],
  "critical_issues": ["...", "..."],
  "improvements": [
    {{
      "area": "Summary",
      "problem": "...",
      "recommendation": "..."
    }}
  ],
  "keyword_analysis": {{
    "matched_keywords": ["..."],
    "missing_or_weak_keywords": ["..."]
  }},
  "formatting_check": {{
    "ats_friendly": true,
    "issues": ["..."]
  }},
  "rewrites": [
    {{
      "original": "quote or short description of the weak bullet/section",
      "improved": "better ATS-friendly version without inventing facts"
    }}
  ],
  "priority_actions": ["...", "...", "..."]
}}

Job description:
{job_context}

Resume:
{resume_text[:30000]}
"""

    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=prompt,
        config={
            "temperature": 0.2,
            "response_mime_type": "application/json",
        },
    )

    raw = response.text.strip()
    try:
        return json.loads(raw)
    except json.JSONDecodeError as exc:
        # Recover if the model accidentally wraps JSON in markdown.
        cleaned = re.sub(r"^```json\s*|\s*```$", "", raw, flags=re.I).strip()
        try:
            return json.loads(cleaned)
        except json.JSONDecodeError:
            raise RuntimeError("Gemini returned an invalid analysis format.") from exc


def show_score(score):
    st.metric("Estimated ATS score", f"{score}/100")
    st.progress(score / 100)


st.title("📄 ATS Resume Checker")
st.write(
    "Upload your resume to get an estimated ATS score, keyword feedback, "
    "formatting checks, and practical improvement suggestions."
)

with st.sidebar:
    st.header("Optional")
    job_description = st.text_area(
        "Paste the job description",
        height=250,
        placeholder="Paste the job posting here for keyword matching...",
    )
    st.caption(
        "The job description is optional. Adding it makes the keyword analysis "
        "more targeted."
    )

uploaded_file = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx"],
    help="PDF or DOCX only. Text-based PDFs work best.",
)

if uploaded_file:
    st.success(f"Uploaded: {uploaded_file.name}")

    if st.button("🔍 Analyze Resume", type="primary", use_container_width=True):
        with st.spinner("Reading and analyzing your resume..."):
            try:
                resume_text = extract_resume_text(uploaded_file)

                if len(resume_text.strip()) < 80:
                    st.error(
                        "Very little text could be extracted. If this is a scanned/image PDF, "
                        "please use a text-based PDF or DOCX."
                    )
                    st.stop()

                baseline_score, baseline_issues = heuristic_ats_score(resume_text)

                result = analyze_with_gemini(resume_text, job_description)

                ai_score = int(result.get("ats_score", baseline_score))
                ai_score = max(0, min(100, ai_score))

                st.divider()
                show_score(ai_score)

                st.caption(
                    f"Local parsing baseline: {baseline_score}/100. "
                    "The final score is Gemini's estimate and should be treated as guidance, "
                    "not a guarantee of passing any particular ATS."
                )

                col1, col2 = st.columns(2)

                with col1:
                    st.subheader("✅ Strengths")
                    for item in result.get("strengths", []):
                        st.write(f"• {item}")

                with col2:
                    st.subheader("⚠️ Critical issues")
                    for item in result.get("critical_issues", []):
                        st.write(f"• {item}")

                st.subheader("🧠 Score explanation")
                st.write(result.get("score_explanation", ""))

                st.subheader("🔑 Keyword analysis")
                keywords = result.get("keyword_analysis", {})
                c1, c2 = st.columns(2)
                with c1:
                    st.write("**Matched keywords**")
                    st.write(", ".join(keywords.get("matched_keywords", [])) or "None identified.")
                with c2:
                    st.write("**Missing / weak keywords**")
                    st.write(", ".join(keywords.get("missing_or_weak_keywords", [])) or "None identified.")

                st.subheader("🛠️ Recommended improvements")
                for item in result.get("improvements", []):
                    with st.expander(item.get("area", "Improvement")):
                        st.write("**Problem:**", item.get("problem", ""))
                        st.write("**Recommendation:**", item.get("recommendation", ""))

                st.subheader("✍️ Suggested rewrites")
                rewrites = result.get("rewrites", [])
                if rewrites:
                    for item in rewrites:
                        st.write("**Original / weak version**")
                        st.info(item.get("original", ""))
                        st.write("**Improved version**")
                        st.success(item.get("improved", ""))
                else:
                    st.write("No specific rewrites were identified.")

                st.subheader("📐 Formatting / ATS check")
                formatting = result.get("formatting_check", {})
                st.write(
                    "ATS-friendly baseline:",
                    "Yes" if formatting.get("ats_friendly") else "Needs work",
                )
                for item in formatting.get("issues", []):
                    st.write(f"• {item}")

                st.subheader("🚀 Top priority actions")
                for i, action in enumerate(result.get("priority_actions", []), 1):
                    st.write(f"{i}. {action}")

                with st.expander("Technical baseline checks"):
                    for issue in baseline_issues:
                        st.write(f"• {issue}")

                with st.expander("Extracted resume text"):
                    st.text(resume_text)

            except Exception as exc:
                st.error(f"Could not analyze the resume: {exc}")
                st.info(
                    "Check that GEMINI_API_KEY is configured and that the uploaded "
                    "resume is a text-based PDF or DOCX."
                )
else:
    st.info("Upload a PDF or DOCX resume to begin.")

st.divider()
st.caption(
    "Privacy note: this demo sends the extracted resume text to Gemini for analysis. "
    "Do not upload documents containing information you are not comfortable sending to "
    "the configured AI service."
)
